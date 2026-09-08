import io
import re
import zipfile
from typing import Dict, List, Optional, Tuple
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

TARGET_COLUMNS = [
    "interest_id",
    "unit_code",
    "security_code",
    "int_type",
    "isin_code",
    "dpid",
    "holder_folio",
    "holder",
    "holder_addr1",
    "holder_pin",
    "bank_accno",
    "bank_name",
    "ifsc_code",
    "ecs_actype",
    "bfitpan",
    "hold_minor",
    "bonds",
    "princ_amt",
    "gross_amt",
    "int_per",
    "face_value",
    "fr_date",
    "to_date",
    "days",
    "wardate",
    "mode_pay",
    "warno",
    "war_acno",
    "type",
    "benpos_date",
]

# Common variations/synonyms for financial/interest registers
SYNONYMS: Dict[str, List[str]] = {
    "interest_id": ["dividend_id", "interestid", "dividendid", "int_id", "intid"],
    "isin_code": ["isin", "isincode", "isin_no", "isin_number", "isin_num"],
    "unit_code": ["unitcode", "unit_cd", "company_code", "comp_code"],
    "security_code": ["securitycode", "security_cd", "sec_code", "seccod", "sec_cod"],
    "int_type": ["inttype", "interest_type", "type_of_interest"],
    "holder_folio": ["folio", "folio_no", "foliono", "holderfolio", "folio_number", "reg_folio"],
    "holder": ["holder_name", "holdername", "investor_name", "first_holder_name", "shareholder_name"],
    "holder_pin": ["holder_pincode", "holderpin", "pincode", "pin_code", "pin"],
    "bank_accno": ["bank_account_no", "bank_acc_no", "bank_account", "bank_ac_no", "bank_acno", "acc_no", "account_no", "acno"],
    "ifsc_code": ["ifsc", "ifsccode", "ifsc_cd"],
    "bfitpan": ["pan", "pan_no", "panno", "pan_number", "first_holder_pan"],
    "ecs_actype": ["ecs_ac_type", "account_type", "ac_type", "actype"],
}


def normalize_col_name(col: str) -> str:
    """Normalize column name for robust case-insensitive matching."""
    if not isinstance(col, str):
        col = str(col) if pd.notna(col) else ""
    normalized = col.strip().lower()
    normalized = re.sub(r"[\s\-]+", "_", normalized)
    return normalized


def detect_header_row(raw_df: pd.DataFrame, max_rows: int = 30) -> Tuple[int, int]:
    """
    Scans the first `max_rows` rows of a raw DataFrame (loaded with header=None).
    Returns:
        (best_row_index, match_count) where best_row_index is 0-indexed.
    """
    target_norms = {normalize_col_name(tc): tc for tc in TARGET_COLUMNS}
    synonym_norms = {}
    for canonical, syn_list in SYNONYMS.items():
        for syn in syn_list:
            synonym_norms[normalize_col_name(syn)] = canonical

    best_row = 0
    best_matches = 0
    limit = min(max_rows, len(raw_df))

    for r_idx in range(limit):
        row_vals = [normalize_col_name(v) for v in raw_df.iloc[r_idx].values if pd.notna(v)]
        matched_targets = set()
        for val in row_vals:
            if val in target_norms:
                matched_targets.add(target_norms[val])
            elif val in synonym_norms:
                matched_targets.add(synonym_norms[val])

        match_count = len(matched_targets)
        if match_count > best_matches:
            best_matches = match_count
            best_row = r_idx

    return best_row, best_matches


def map_columns(source_columns: List[str]) -> Tuple[Dict[str, str], List[str], List[str]]:
    """
    Maps source dataframe column names to canonical target column names.
    Returns:
        col_mapping: dict of {source_col_name: canonical_target_col_name}
        missing_targets: list of canonical target columns not found in source
        dropped_columns: list of source columns that will be discarded
    """
    target_norm_map = {normalize_col_name(tc): tc for tc in TARGET_COLUMNS}
    synonym_norm_map = {}
    for canonical, syn_list in SYNONYMS.items():
        for syn in syn_list:
            synonym_norm_map[normalize_col_name(syn)] = canonical

    col_mapping = {}
    found_targets = set()
    dropped_columns = []

    for src_col in source_columns:
        norm_src = normalize_col_name(src_col)
        if norm_src in target_norm_map:
            canonical = target_norm_map[norm_src]
            col_mapping[src_col] = canonical
            found_targets.add(canonical)
        elif norm_src in synonym_norm_map:
            canonical = synonym_norm_map[norm_src]
            col_mapping[src_col] = canonical
            found_targets.add(canonical)
        else:
            dropped_columns.append(src_col)

    missing_targets = [tc for tc in TARGET_COLUMNS if tc not in found_targets]
    return col_mapping, missing_targets, dropped_columns


def process_dataframe(
    df: pd.DataFrame, fill_missing_cols: bool = True, drop_summary_rows: bool = True
) -> Tuple[pd.DataFrame, Dict[str, str], List[str], List[str], int]:
    """
    Filters df to only the target columns in the exact order specified.
    Renames matched columns to canonical names.
    Optionally drops blank / summary / grand total footer rows.
    """
    col_mapping, missing_targets, dropped_columns = map_columns(list(df.columns))

    # Rename matched columns to canonical target names
    filtered_df = df.rename(columns=col_mapping)

    # Keep only matched target columns
    cols_to_keep = [col for col in TARGET_COLUMNS if col in filtered_df.columns]
    filtered_df = filtered_df[cols_to_keep].copy()

    # Fill missing columns with None if requested to preserve strict 30-column layout
    if fill_missing_cols:
        for missing_col in missing_targets:
            filtered_df[missing_col] = None
        filtered_df = filtered_df[TARGET_COLUMNS]

    dropped_summary_rows = 0
    if drop_summary_rows and "isin_code" in filtered_df.columns:
        initial_len = len(filtered_df)

        def is_valid_data_row(row) -> bool:
            isin_val = str(row["isin_code"]).strip() if pd.notna(row["isin_code"]) else ""
            # If isin_code is present and doesn't say "total", it is a valid row
            if isin_val and not isin_val.lower().startswith("total"):
                return True
            # If isin is missing, check if this is a grand total / summary footer row
            # e.g., holder and interest_id are both missing, or contain "total"
            holder_val = str(row.get("holder", "")).strip().lower() if pd.notna(row.get("holder")) else ""
            iid_val = str(row.get("interest_id", "")).strip().lower() if pd.notna(row.get("interest_id")) else ""
            folio_val = str(row.get("holder_folio", "")).strip().lower() if pd.notna(row.get("holder_folio")) else ""

            if not holder_val and not iid_val and not folio_val:
                return False  # Blank row or grand total footer
            if "total" in holder_val or "total" in iid_val or "total" in folio_val:
                return False  # Total summary row
            return True

        mask = filtered_df.apply(is_valid_data_row, axis=1)
        filtered_df = filtered_df[mask].copy()
        dropped_summary_rows = initial_len - len(filtered_df)

    return filtered_df, col_mapping, missing_targets, dropped_columns, dropped_summary_rows


def split_by_isin(df: pd.DataFrame) -> Dict[str, pd.DataFrame]:
    """
    Splits the dataframe into groups keyed by isin_code.
    Rows with blank/null isin_code are grouped under 'UNASSIGNED_ISIN'.
    """
    if "isin_code" not in df.columns:
        raise ValueError("Column 'isin_code' is missing from the data.")

    # Create cleaned isin series
    isin_series = df["isin_code"].fillna("").astype(str).str.strip()

    groups: Dict[str, pd.DataFrame] = {}
    unique_isins = isin_series.unique()

    for isin in unique_isins:
        key = isin if isin != "" else "UNASSIGNED_ISIN"
        mask = (isin_series == isin)
        sub_df = df[mask].copy()
        groups[key] = sub_df

    return groups


def dataframe_to_styled_excel(df: pd.DataFrame, sheet_name: str = "Sheet1") -> bytes:
    """
    Converts a DataFrame to a professionally formatted Excel (.xlsx) file in memory.
    Applies styling: bold header, themed header fill, thin borders, and auto column widths.
    """
    output = io.BytesIO()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    # Styles
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    data_font = Font(name="Segoe UI", size=9)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    # Write headers
    headers = list(df.columns)
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Write data rows
    for row_num, row_data in enumerate(df.itertuples(index=False), 2):
        for col_num, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_num, value=val)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = left_align

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value if cell.value is not None else "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Freeze top row
    ws.freeze_panes = "A2"

    wb.save(output)
    output.seek(0)
    return output.getvalue()


def create_isin_zip_archive(isin_groups: Dict[str, pd.DataFrame]) -> bytes:
    """
    Creates an in-memory ZIP file containing an Excel file for each ISIN group.
    """
    zip_buffer = io.BytesIO()

    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
        for isin, sub_df in isin_groups.items():
            # Sanitize filename
            safe_isin = re.sub(r'[\\/*?:"<>|]', "_", str(isin))
            filename = f"{safe_isin}.xlsx"
            excel_bytes = dataframe_to_styled_excel(sub_df, sheet_name=safe_isin[:31])
            zip_file.writestr(filename, excel_bytes)

    zip_buffer.seek(0)
    return zip_buffer.getvalue()
